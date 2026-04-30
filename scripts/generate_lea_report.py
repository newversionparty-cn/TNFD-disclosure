#!/usr/bin/env python3
"""
Generate a TNFD LEA preliminary report from saved project state.
Exports saved state only; does not infer new conclusions.
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_STATE = Path(os.environ.get("TNFD_HOME", SKILL_DIR / ".tnfd")) / "project-state.json"
DEFAULT_CONFIG = Path(os.environ.get("TNFD_HOME", SKILL_DIR / ".tnfd")) / "config.json"
DEFAULT_OUT_DIR = SKILL_DIR / "outputs" / "tnfd_lea"


def load_json(path: Path, fallback: Any) -> Any:
    if path.exists():
        with path.open(encoding="utf-8") as f:
            return json.load(f)
    return fallback


def current_project(state: dict, config: dict) -> dict:
    projects = state.get("projects", {})
    current_id = config.get("current_project")
    if current_id and current_id in projects:
        project = projects[current_id]
        project.setdefault("id", current_id)
        return project
    if projects:
        project_id, project = next(iter(projects.items()))
        project.setdefault("id", project_id)
        return project
    return {
        "id": "no_project",
        "name": "Evidence gap",
        "industry": "Evidence gap",
        "mode": "beginner_lea",
        "role_frame": "big4_pua_senior",
        "partner_pressure_level": "L2",
        "lea_outputs": {},
        "risks_found": [],
        "artifacts": {}
    }


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    if not rows:
        rows = [["Evidence gap" for _ in headers]]
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        padded = list(row) + [""] * (len(headers) - len(row))
        lines.append("| " + " | ".join(str(v) if v is not None else "" for v in padded[:len(headers)]) + " |")
    return "\n".join(lines)


def rows_from(items: Any, columns: list[str]) -> list[list[Any]]:
    if not items:
        return [["Evidence gap" for _ in columns]]
    if isinstance(items, dict):
        items = [items]
    rows = []
    for item in items:
        if isinstance(item, dict):
            rows.append([item.get(col, item.get(col.lower(), "Evidence gap")) for col in columns])
        else:
            rows.append([str(item)] + ["Evidence gap"] * (len(columns) - 1))
    return rows


def report_markdown(project: dict) -> str:
    outputs = project.get("lea_outputs", {})
    locate = outputs.get("locate", {})
    evaluate = outputs.get("evaluate", {})
    assess = outputs.get("assess", {})

    asset_cols = ["asset", "type", "address_or_city", "data_precision", "nature_interface", "initial_sensitivity", "gap"]
    matrix_cols = ["activity", "dependency", "impact", "importance", "evidence_status", "confidence"]
    risk_cols = ["risk_or_opportunity", "type", "trigger", "business_impact", "financial_pathway", "priority", "required_evidence"]
    gap_cols = ["gap", "affects", "severity", "owner", "next_action"]
    gaps = []
    for stage in outputs.values():
        if isinstance(stage, dict):
            gaps.extend(stage.get("evidence_gaps", []))

    return f"""# TNFD LEA Preliminary Report

## Executive Summary

合伙人，先给结论：本报告仅基于已保存 LEA 状态导出，不新增推断。当前项目为 `{project.get('name', 'Evidence gap')}`，行业为 `{project.get('industry', 'Evidence gap')}`，Partner Pressure Level 为 `{project.get('partner_pressure_level', 'L2')}`。

## Project Scope

{md_table(['Field', 'Value'], [
    ['Project ID', project.get('id', 'Evidence gap')],
    ['Company', project.get('name', 'Evidence gap')],
    ['Industry', project.get('industry', 'Evidence gap')],
    ['Mode', project.get('mode', 'beginner_lea')],
    ['Role frame', project.get('role_frame', 'big4_pua_senior')],
    ['Exported at', datetime.now().isoformat(timespec='seconds')]
])}

## Locate Findings

{md_table(asset_cols, rows_from(locate.get('assets'), asset_cols))}

## Evaluate Findings

{md_table(matrix_cols, rows_from(evaluate.get('dependency_impact_matrix'), matrix_cols))}

## Assess Findings

{md_table(risk_cols, rows_from(assess.get('risk_opportunity_register'), risk_cols))}

## Evidence Gaps

{md_table(gap_cols, rows_from(gaps, gap_cols))}

## Partner Pressure Review

{md_table(['Level', 'Reason', 'Required rework'], [[
    project.get('partner_pressure_level', 'L2'),
    'Saved LEA state contains evidence gaps or unverified fields',
    'Complete missing evidence before Prepare or external disclosure'
]])}

## Limitations

- This report is an internal preliminary LEA workpaper.
- It is not a full TNFD disclosure.
- It is not a formal audit or assurance opinion.
- Pending case claims remain pending unless verified in the claim registry.
- Internal scoring or screening outputs are not official TNFD risk scores.
"""


def write_docx(markdown_text: str, path: Path) -> bool:
    try:
        from docx import Document
    except Exception:
        return False

    doc = Document()
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            doc.add_paragraph(line)
        elif line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    doc.save(path)
    return True


def write_pdf(markdown_text: str, path: Path) -> bool:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except Exception:
        return False

    doc = SimpleDocTemplate(str(path), pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    for line in markdown_text.splitlines():
        if not line.strip():
            story.append(Spacer(1, 8))
            continue
        style = styles["Heading1"] if line.startswith("# ") else styles["Heading2"] if line.startswith("## ") else styles["BodyText"]
        clean = line.lstrip("# ").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(clean, style))
    doc.build(story)
    return True


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state", type=Path, default=DEFAULT_STATE)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--format", choices=["all", "docx", "pdf", "md"], default="all")
    args = parser.parse_args()

    state = load_json(args.state, {"projects": {}})
    config = load_json(args.config, {})
    project = current_project(state, config)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    markdown_text = report_markdown(project)
    md_path = args.output_dir / "TNFD_LEA_Preliminary_Report.md"
    docx_path = args.output_dir / "TNFD_LEA_Preliminary_Report.docx"
    pdf_path = args.output_dir / "TNFD_LEA_Preliminary_Report.pdf"

    written = []
    if args.format in ["all", "md"]:
        md_path.write_text(markdown_text, encoding="utf-8")
        written.append(md_path)
    if args.format in ["all", "docx"]:
        if write_docx(markdown_text, docx_path):
            written.append(docx_path)
        elif args.format == "docx":
            md_path.write_text(markdown_text, encoding="utf-8")
            written.append(md_path)
    if args.format in ["all", "pdf"]:
        if write_pdf(markdown_text, pdf_path):
            written.append(pdf_path)
        elif args.format == "pdf":
            md_path.write_text(markdown_text, encoding="utf-8")
            written.append(md_path)

    for path in written:
        print(path)
    if not written:
        md_path.write_text(markdown_text, encoding="utf-8")
        print(md_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""
Generate a TNFD LEA Excel workpaper from saved project state.
This script exports saved state only; it does not infer new conclusions.
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_STATE = Path(os.environ.get("TNFD_HOME", SKILL_DIR / ".tnfd")) / "project-state.json"
DEFAULT_CONFIG = Path(os.environ.get("TNFD_HOME", SKILL_DIR / ".tnfd")) / "config.json"
DEFAULT_OUT = SKILL_DIR / "outputs" / "tnfd_lea" / "TNFD_LEA_Workbook.xlsx"


def load_json(path: Path, fallback: Any) -> Any:
    if path.exists():
        with path.open(encoding="utf-8") as f:
            return json.load(f)
    return fallback


def current_project(state: dict, config: dict) -> dict:
    projects = state.get("projects", {})
    current_id = config.get("current_project")
    if current_id and current_id in projects:
        project = projects[current_id]
        project.setdefault("id", current_id)
        return project
    if projects:
        project_id, project = next(iter(projects.items()))
        project.setdefault("id", project_id)
        return project
    return {
        "id": "no_project",
        "name": "Evidence gap",
        "industry": "Evidence gap",
        "mode": "beginner_lea",
        "role_frame": "big4_pua_senior",
        "partner_pressure_level": "L2",
        "lea_outputs": {},
        "risks_found": [],
        "artifacts": {}
    }


def rows_from(items: Any, columns: list[str]) -> list[list[Any]]:
    if not items:
        return [["Evidence gap" for _ in columns]]
    if isinstance(items, dict):
        items = [items]
    rows = []
    for item in items:
        if isinstance(item, dict):
            rows.append([item.get(col, item.get(col.lower(), "Evidence gap")) for col in columns])
        else:
            rows.append([str(item)] + ["Evidence gap"] * (len(columns) - 1))
    return rows


def add_table(ws, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13)
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.append([])


def autosize(ws) -> None:
    for column_cells in ws.columns:
        max_len = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(column)].width = max(14, max_len + 2)


def build_workbook(project: dict) -> Workbook:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    outputs = project.get("lea_outputs", {})

    intake = outputs.get("intake", {})
    locate = outputs.get("locate", {})
    evaluate = outputs.get("evaluate", {})
    assess = outputs.get("assess", {})

    ws = wb.create_sheet("00_Project_Intake")
    add_table(ws, "Project Intake", ["Field", "Value"], [
        ["Project ID", project.get("id", "Evidence gap")],
        ["Company", project.get("name") or intake.get("company", "Evidence gap")],
        ["Industry", project.get("industry") or intake.get("sector", "Evidence gap")],
        ["Mode", project.get("mode", "beginner_lea")],
        ["Role frame", project.get("role_frame", "big4_pua_senior")],
        ["Partner pressure level", project.get("partner_pressure_level", "L2")],
        ["Data maturity", intake.get("data_maturity", "Evidence gap")],
        ["Exported at", datetime.now().isoformat(timespec="seconds")]
    ])
    autosize(ws)

    ws = wb.create_sheet("01_Asset_Register")
    asset_columns = ["asset", "type", "address_or_city", "data_precision", "nature_interface", "initial_sensitivity", "gap"]
    add_table(ws, "Asset Register", asset_columns, rows_from(locate.get("assets"), asset_columns))
    autosize(ws)

    ws = wb.create_sheet("02_Locate_Screening")
    screening_columns = ["asset", "water_related", "land_related", "biodiversity_related", "community_or_regulatory_sensitivity", "priority"]
    add_table(ws, "Locate Screening", screening_columns, rows_from(locate.get("priority_location_screen"), screening_columns))
    autosize(ws)

    ws = wb.create_sheet("03_Dependency_Impact")
    matrix_columns = ["activity", "dependency", "impact", "importance", "evidence_status", "confidence"]
    add_table(ws, "Dependency And Impact Matrix", matrix_columns, rows_from(evaluate.get("dependency_impact_matrix"), matrix_columns))
    autosize(ws)

    ws = wb.create_sheet("04_Risk_Opportunity_Register")
    register_columns = ["risk_or_opportunity", "type", "trigger", "business_impact", "financial_pathway", "priority", "required_evidence"]
    add_table(ws, "Risk And Opportunity Register", register_columns, rows_from(assess.get("risk_opportunity_register"), register_columns))
    autosize(ws)

    ws = wb.create_sheet("05_Evidence_Gaps")
    gap_columns = ["gap", "affects", "severity", "owner", "next_action"]
    gaps = []
    for stage in (intake, locate, evaluate, assess):
        gaps.extend(stage.get("evidence_gaps", []) if isinstance(stage, dict) else [])
    add_table(ws, "Evidence Gaps", gap_columns, rows_from(gaps, gap_columns))
    autosize(ws)

    ws = wb.create_sheet("Dashboard")
    phase_status = project.get("phase_status", {})
    add_table(ws, "LEA Dashboard", ["Metric", "Value"], [
        ["Locate status", phase_status.get("locate", {}).get("status", "not_started")],
        ["Evaluate status", phase_status.get("evaluate", {}).get("status", "not_started")],
        ["Assess status", phase_status.get("assess", {}).get("status", "not_started")],
        ["Risks found", len(project.get("risks_found", []))],
        ["Data quality", project.get("data_quality", "Evidence gap")],
        ["Partner pressure level", project.get("partner_pressure_level", "L2")],
        ["Next action", "Complete evidence gaps before Prepare"]
    ])
    autosize(ws)
    return wb


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state", type=Path, default=DEFAULT_STATE)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    state = load_json(args.state, {"projects": {}})
    config = load_json(args.config, {})
    project = current_project(state, config)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(project)
    wb.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
